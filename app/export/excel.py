from pathlib import Path
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from app.database.models import Employee, Shift, AttendanceSession, ScanEvent, ExceptionRecord
from app.attendance.rotation import shift_id_for

SHEET_NAMES = ['Attendance Summary', 'Raw Scans', 'Exceptions', 'Employees', 'Shifts']

class ExcelExporter:
    def __init__(self, db): self.db = db
    def export(self, directory='exports', date_from: date | None = None, date_to: date | None = None,
               employee_id: str | None = None, department: str | None = None, include=None):
        """Write an .xlsx workbook. `include` may restrict which sheets are produced."""
        out = Path(directory); out.mkdir(parents=True, exist_ok=True)
        path = out / f'Attendance_{datetime.now():%Y-%m-%d}.xlsx'
        wanted = list(include) if include else SHEET_NAMES
        wb = Workbook(); wb.remove(wb.active)
        sheets = {name: wb.create_sheet(name) for name in wanted}
        with self.db.session() as s:
            if 'Attendance Summary' in sheets:
                rows = []
                for x in s.scalars(select(AttendanceSession)):
                    e = s.get(Employee, x.employee_id); sh = s.get(Shift, x.shift_id) if x.shift_id else None
                    if (date_from and x.clock_in.date() < date_from) or (date_to and x.clock_in.date() > date_to) or (employee_id and e.employee_id != employee_id) or (department and e.department != department): continue
                    worked = (x.clock_out - x.clock_in) if x.clock_out else None
                    rows.append([e.employee_id, f'{e.first_name} {e.last_name}', x.clock_in.date(), sh.name if sh else '', x.clock_in, x.clock_out, round(worked.total_seconds()/3600, 2) if worked else None, 'OK' if x.clock_out else 'MISSING OUT'])
                self._write(sheets['Attendance Summary'], ['Employee ID', 'Employee Name', 'Date', 'Shift', 'Clock In', 'Clock Out', 'Worked Hours', 'Status'], rows)
            if 'Raw Scans' in sheets:
                raw = []
                for x in s.scalars(select(ScanEvent).order_by(ScanEvent.timestamp)):
                    e = s.get(Employee, x.employee_id) if x.employee_id else None
                    if (date_from and x.timestamp.date() < date_from) or (date_to and x.timestamp.date() > date_to) or (employee_id and x.employee_id != employee_id) or (department and (not e or e.department != department)): continue
                    raw.append([x.timestamp, x.employee_id or '', f'{e.first_name} {e.last_name}' if e else '', x.action, x.accepted, x.terminal_id or '', x.scanner_id or '', x.raw_barcode, x.rejection_reason or ''])
                self._write(sheets['Raw Scans'], ['Timestamp', 'Employee ID', 'Employee Name', 'Action', 'Accepted', 'Terminal', 'Scanner', 'Raw Barcode', 'Reason'], raw)
            if 'Exceptions' in sheets:
                self._write(sheets['Exceptions'], ['Timestamp', 'Employee', 'Problem', 'Terminal', 'Scanner', 'Details'], [[x.timestamp, x.employee_id or '', x.problem, x.terminal_id or '', x.scanner_id or '', x.details or ''] for x in s.scalars(select(ExceptionRecord))])
            if 'Employees' in sheets:
                shift_names = {x.id: x.name for x in s.scalars(select(Shift))}
                employee_rows = []
                for e in s.scalars(select(Employee)):
                    shift_id = shift_id_for(s, e, date.today())
                    shift_name = shift_names.get(shift_id, '') if shift_id else ''
                    employee_rows.append([e.employee_id, e.employee_number or '', e.first_name, e.last_name,
                                          e.department or '', e.position or '', e.active, shift_name])
                self._write(sheets['Employees'], ['Employee ID', 'Number', 'First Name', 'Last Name', 'Department', 'Position', 'Active', 'Shift'], employee_rows)
            if 'Shifts' in sheets:
                self._write(sheets['Shifts'], ['Name', 'Start', 'End', 'Early clock-in min', 'Earliest clock-out min', 'Active'], [[x.name, x.start_time, x.end_time, x.early_clock_in_minutes, x.earliest_clock_out_minutes, x.active] for x in s.scalars(select(Shift))])
        wb.save(path); return path
    def _write(self, ws, headers, rows):
        ws.append(headers)
        for row in rows: ws.append(row)
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
        for i in range(1, len(headers)+1): ws.column_dimensions[get_column_letter(i)].width = max(14, min(35, len(headers[i-1])+4))
