"""Employee list import from an Excel (.xlsx) or CSV file, plus template creation.

Accepted columns (headers are matched case-insensitively, ignoring spaces/underscores):

    first_name | last_name | employee_number | department | position | shift

* ``first_name`` and ``last_name`` are required for a row to be imported.
* ``shift`` is a shift *name*; it must match an existing shift or the row is
  imported without a shift (a note is returned).
* ``employee_number`` is optional; when present, rows whose number already
  exists are skipped as duplicates.
"""
from __future__ import annotations
from pathlib import Path

HEADERS = ['first_name', 'last_name', 'employee_number', 'department', 'position', 'shift']
REQUIRED = ('first_name', 'last_name')

EXAMPLE_ROWS = [
    {'first_name': 'Maria', 'last_name': 'Pop', 'employee_number': '1001',
     'department': 'Sales', 'position': 'Sales Representative', 'shift': ''},
    {'first_name': 'Ion', 'last_name': 'Popescu', 'employee_number': '1002',
     'department': 'IT', 'position': 'Support Technician', 'shift': ''},
]


def _normalise_header(value) -> str:
    """'First Name' -> 'first_name', 'employeenumber' -> 'employee_number'."""
    compact = ''.join(ch for ch in str(value).lower() if ch.isalnum())
    mapping = {'firstname': 'first_name', 'lastname': 'last_name',
               'employeenumber': 'employee_number', 'department': 'department',
               'position': 'position', 'shift': 'shift', 'shiftname': 'shift'}
    return mapping.get(compact, compact)


def _clean_row(mapping: dict) -> dict:
    row = {key: '' for key in HEADERS}
    for header, value in mapping.items():
        key = _normalise_header(header)
        if key in row:
            row[key] = ('' if value is None else str(value)).strip()
    return row


def read_employee_file(path: str | Path) -> list[dict]:
    """Read an .xlsx or .csv file into a list of row dicts (headers normalised).

    The first row is always treated as the header.
    """
    path = Path(path)
    if not path.exists():
        raise ValueError(f'File not found: {path}')
    raw = _raw_rows(path)
    if not raw:
        raise ValueError('The file is empty.')
    headers = [str(header).strip() for header in raw[0]]
    result = []
    for line in raw[1:]:
        mapping = {}
        for index, header in enumerate(headers):
            value = line[index] if index < len(line) else None
            mapping[header] = '' if value is None else str(value).strip()
        row = _clean_row(mapping)
        if not any(row.values()):      # skip fully empty lines
            continue
        result.append(row)
    if not result:
        raise ValueError('The file contains no employee rows.')
    return result


def _raw_rows(path: Path) -> list[list]:
    suffix = path.suffix.lower()
    if suffix == '.csv':
        import csv
        with open(path, 'r', encoding='utf-8-sig', newline='') as handle:
            return [list(line) for line in csv.reader(handle)]
    if suffix in ('.xlsx', '.xlsm'):
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise RuntimeError('openpyxl is required to read Excel files.') from error
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = [['' if cell is None else cell for cell in row]
                for row in workbook.active.iter_rows(values_only=True)]
        workbook.close()
        return rows
    raise ValueError('Unsupported file type. Use an .xlsx or .csv file.')


# -- templates ---------------------------------------------------------------
def write_template_xlsx(path: str | Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as error:
        raise RuntimeError('openpyxl is required to create the Excel template.') from error
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Employees'
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F4E78')
    for example in EXAMPLE_ROWS:
        sheet.append([example.get(key, '') for key in HEADERS])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, len(HEADERS) + 1):
        letter = sheet.cell(row=1, column=column).column_letter
        sheet.column_dimensions[letter].width = 20
    workbook.save(path)
    return path


def write_template_csv(path: str | Path) -> Path:
    import csv
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        for example in EXAMPLE_ROWS:
            writer.writerow([example.get(key, '') for key in HEADERS])
    return path
