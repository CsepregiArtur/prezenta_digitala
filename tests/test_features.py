"""Automated checks for the features added to complete the implementation phases:
hardware CRUD, scheduled export job frequency/next-run, Excel sheet selection,
dashboard summary semantics, LAN/server synchronization, database migration, and
the threaded scanner manager pipeline.
"""
import json
import sqlite3
import sys
import threading
import time as _time
import types
from datetime import datetime, time, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest
from app.database import Database
from app.database.models import AttendanceSession, Shift
from app.services import EmployeeService, ShiftService, AdminDataService
from app.attendance import AttendanceEngine, ScannerEvent
from app.scheduler import next_run_time, is_due, schedule_text


@pytest.fixture
def db(tmp_path):
    d = Database(tmp_path / 'attendance.db'); d.seed_defaults()
    with d.session() as s:
        s.add_all([Shift(name='Day', start_time=time(6), end_time=time(14), early_clock_in_minutes=60, earliest_clock_out_minutes=10)])
        s.commit()
    return d


def employee(db, shift=1):
    return EmployeeService(db).create('Maria', 'Pop', shift_id=shift)


# -- Phase 6: terminal/scanner edit, disable and delete ------------------------
def test_scanner_terminal_crud(db):
    service = AdminDataService(db)
    service.save_scanner('S1', 1, 'COM3', 'Front door', baud_rate=115200)
    assert service.get_scanner('S1').baud_rate == 115200
    service.save_scanner('S2', 1, 'COM5')
    assert len(service.scanners_flat()) == 2
    service.update_scanner('S1', com_port='COM4', enabled=False)
    assert service.get_scanner('S1').com_port == 'COM4' and not service.get_scanner('S1').enabled
    service.delete_scanner('S1')
    assert service.get_scanner('S1') is None and len(service.scanners_flat()) == 1
    service.update_terminal(1, name='Main gate')
    assert service.terminals()[0][0].name == 'Main gate'
    service.delete_terminal(1)
    # deleting a terminal removes its scanners (FK clean-up)
    assert len(service.terminals()) == 1 and service.terminals()[0][0].name == 'Terminal 2'
    assert service.scanners_flat() == []


# -- Phase 7: dashboard summary card semantics ---------------------------------
def test_summary_present_vs_missing_clockout(db):
    eid = employee(db)
    today = datetime.now().date()
    with db.session() as s:
        s.add(AttendanceSession(employee_id=eid, clock_in=datetime(today.year, today.month, today.day, 7)))
        s.add(AttendanceSession(employee_id=eid, clock_in=datetime(today.year, today.month, today.day - 1, 22)))
        s.commit()
    summary = AdminDataService(db).summary()
    assert summary['present_now'] == 1
    assert summary['missing_out'] == 1
    assert summary['open_sessions'] == 2
    assert summary['terminals'] == 2


# -- Phase 10: export job schedule fields and next-run -------------------------
def test_export_job_frequency_fields(db):
    service = AdminDataService(db)
    service.save_export_job('07:30', 'exports', 'weekly', weekday=0)
    service.save_export_job('08:00', 'exports', 'monthly', month_day=15)
    weekly = next(j for j in service.export_jobs() if j.frequency == 'weekly')
    monthly = next(j for j in service.export_jobs() if j.frequency == 'monthly')
    assert weekly.weekday == 0 and monthly.month_day == 15
    assert schedule_text(weekly) == 'Monday' and schedule_text(monthly) == 'Day 15'
    now = datetime(2026, 9, 1, 9, 0)
    nxt = next_run_time(weekly, now)
    assert nxt.strftime('%H:%M') == '07:30' and nxt.weekday() == 0 and nxt > now
    monthly_next = next_run_time(monthly, datetime(2026, 1, 16, 0, 0))
    assert monthly_next.day == 15 and monthly_next.strftime('%H:%M') == '08:00'
    service.update_export_job(weekly.id, enabled=False)
    assert not next(j for j in service.export_jobs() if j.id == weekly.id).enabled
    service.delete_export_job(monthly.id)
    assert all(j.frequency != 'monthly' for j in service.export_jobs())


def test_is_due_frequency(db):
    service = AdminDataService(db)
    service.save_export_job('09:00', 'x', 'weekly', weekday=2)
    service.save_export_job('09:00', 'x', 'monthly', month_day=15)
    weekly = next(j for j in service.export_jobs() if j.frequency == 'weekly')
    monthly = next(j for j in service.export_jobs() if j.frequency == 'monthly')
    base = datetime(2026, 9, 1, 9, 0)
    while base.weekday() != weekly.weekday: base += timedelta(days=1)
    assert is_due(weekly, base) is True
    assert is_due(weekly, base + timedelta(days=1)) is False
    assert is_due(monthly, datetime(2026, 9, 15, 9, 0)) is True
    assert is_due(monthly, datetime(2026, 9, 14, 9, 0)) is False


# -- Phase 9: Excel sheet selection --------------------------------------------
def test_excel_export_sheet_selection(db, tmp_path):
    pytest.importorskip('openpyxl')
    from app.export import ExcelExporter
    from openpyxl import load_workbook
    selected = ExcelExporter(db).export(tmp_path, include=['Employees', 'Shifts'])
    assert load_workbook(selected).sheetnames == ['Employees', 'Shifts']
    full = ExcelExporter(db).export(tmp_path / 'full')
    assert len(load_workbook(full).sheetnames) == 5


# -- Phase 12: LAN/server synchronization ---------------------------------------
@pytest.fixture
def capture_server():
    received = []
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            length = int(self.headers.get('Content-Length', 0))
            received.append(json.loads(self.rfile.read(length).decode('utf-8')))
            self.send_response(200); self.end_headers(); self.wfile.write(b'{}')
        def log_message(self, *args): pass
    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
    yield f'http://127.0.0.1:{server.server_address[1]}/api/scans', received
    server.shutdown()


def test_sync_pushes_events_once(db, capture_server):
    from app.sync import SyncService
    url, received = capture_server
    eid = employee(db)
    engine = AttendanceEngine(db)
    engine.process(ScannerEvent(eid, datetime(2025, 1, 2, 6), 1, 'S1'))
    engine.process(ScannerEvent(eid, datetime(2025, 1, 2, 14), 1, 'S1'))
    sync = SyncService(db, engine, None)
    assert sync.push_pending(url) == 2
    assert len(received) == 1 and len(received[0]['events']) == 2
    assert received[0]['events'][0]['employee_id'] == eid
    # Nothing new to send the second time.
    assert sync.push_pending(url) == 0 and len(received) == 1
    ok, message = sync.test_connection(url)
    assert ok and 'HTTP' in message
    # A later event is pushed on the next run.
    engine.process(ScannerEvent(eid, datetime(2025, 1, 3, 6), 1, 'S1'))
    assert sync.push_pending(url) == 1


def test_sync_drains_offline_queue(db, tmp_path):
    from app.offline import OfflineQueue
    from app.sync import SyncService
    queue = OfflineQueue(tmp_path / 'queue.db')
    eid = employee(db)
    engine = AttendanceEngine(db)
    queue.enqueue(ScannerEvent(eid, datetime(2025, 1, 2, 6), 1, 'S1'))
    sync = SyncService(db, engine, queue)
    done, failed = sync.drain_offline()
    assert (done, failed, queue.count()) == (1, 0, 0)


# -- Database migration: add columns to an existing file ------------------------
def test_database_migrates_new_columns(tmp_path):
    path = tmp_path / 'existing.db'
    first = Database(path); first.seed_defaults(); first.engine.dispose()
    con = sqlite3.connect(path)
    con.execute('ALTER TABLE export_jobs DROP COLUMN weekday')
    con.execute('ALTER TABLE export_jobs DROP COLUMN month_day')
    con.commit(); con.close()
    Database(path)
    con = sqlite3.connect(path)
    columns = [row[1] for row in con.execute('PRAGMA table_info(export_jobs)')]
    con.close()
    assert 'weekday' in columns and 'month_day' in columns


# -- Phase 5: threaded scanner manager pipeline ---------------------------------
def test_scanner_manager_connects_and_processes(db, monkeypatch):
    from types import SimpleNamespace
    from app.scanner import ScannerManager
    eid = employee(db)
    engine = AttendanceEngine(db)

    class FakePort:
        def __init__(self, *args, **kwargs):
            self.lines = [b'EMP000001\r\n']; self.closed = False
        def __enter__(self): return self
        def __exit__(self, *args): self.close()
        def close(self): self.closed = True
        def readline(self):
            if not self.lines:
                _time.sleep(0.05); raise OSError('simulated disconnect')
            return self.lines.pop(0)

    fake = types.ModuleType('serial'); fake.Serial = FakePort
    monkeypatch.setitem(sys.modules, 'serial', fake)

    manager = ScannerManager(engine)
    scanner = SimpleNamespace(scanner_id='S1', com_port='COM3', baud_rate=9600, enabled=True, terminal_id=1)
    manager.configure([scanner]).start()
    deadline = _time.time() + 5
    while _time.time() < deadline and manager.last_result is None:
        _time.sleep(0.05)
    manager.stop()
    assert manager.last_result is not None and manager.last_result.accepted
    assert manager.status.get('S1') in ('connected', 'offline')
    with db.session() as s:
        from sqlalchemy import select, func
        from app.database.models import ScanEvent
        assert s.scalar(select(func.count()).select_from(ScanEvent)) == 1
