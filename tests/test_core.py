from datetime import datetime, time
from pathlib import Path
import pytest
from app.database import Database
from app.database.models import Employee, Shift, Terminal, Scanner, ScanEvent, AttendanceSession
from app.services import EmployeeService
from app.services import ShiftService, AdminDataService
from app.attendance import AttendanceEngine, ScannerEvent
from app.auth import AuthService
from app.maintenance import backup_database

@pytest.fixture
def db(tmp_path):
    d=Database(tmp_path/'attendance.db'); d.seed_defaults()
    with d.session() as s:
        s.add_all([Shift(name='Day',start_time=time(6),end_time=time(14),early_clock_in_minutes=60,earliest_clock_out_minutes=10),Shift(name='Night',start_time=time(22),end_time=time(6),early_clock_in_minutes=60,earliest_clock_out_minutes=10)]); s.commit()
    return d
def employee(db, shift=1): return EmployeeService(db).create('Maria','Pop',shift_id=shift)
def scan(engine,eid,dt,scanner='S1'): return engine.process(ScannerEvent(eid,dt,1,scanner))
def test_employee_creation(db):
    eid=employee(db); assert eid=='EMP000001'
    with db.session() as s: assert s.get(Employee,eid).last_name=='Pop'
def test_shift_creation_and_editing(db):
    service=ShiftService(db); service.create('Evening',time(14),time(22),30,15)
    shift=next(x for x in service.list() if x.name=='Evening')
    service.update(shift.id,early_clock_in_minutes=45,active=False)
    updated=next(x for x in service.list() if x.id==shift.id)
    assert updated.early_clock_in_minutes==45 and not updated.active
def test_scanner_configuration(db):
    service=AdminDataService(db); service.save_scanner('Scanner 1',1,'COM3','Front door')
    service.update_scanner('Scanner 1',com_port='COM4',enabled=False)
    terminal, scanners=service.terminals()[0]
    assert terminal.name=='Terminal 1' and scanners[0].com_port=='COM4' and not scanners[0].enabled
def test_scheduled_export_configuration(db,tmp_path):
    service=AdminDataService(db); service.save_export_job('10:00',str(tmp_path))
    job=service.export_jobs()[0]; service.update_export_job(job.id,enabled=False)
    assert service.export_jobs()[0].run_time=='10:00' and not service.export_jobs()[0].enabled
def test_offline_queue_synchronizes_without_losing_event(db,tmp_path):
    from app.offline import OfflineQueue
    queue=OfflineQueue(tmp_path/'offline.db'); eid=employee(db); event=ScannerEvent(eid,datetime(2025,1,2,6),1,'SIM')
    queue.enqueue(event); assert queue.count()==1
    done,failed=queue.synchronize(AttendanceEngine(db))
    assert (done,failed,queue.count())==(1,0,0)
    with db.session() as s: assert s.scalar(__import__('sqlalchemy').select(__import__('sqlalchemy').func.count()).select_from(ScanEvent))==1

def test_end_to_end_day_shift_with_rejected_out_and_excel(db,tmp_path):
    """Master specification end-to-end scenario: accepted IN, rejected OUT, accepted OUT."""
    eid=employee(db); engine=AttendanceEngine(db)
    assert scan(engine,eid,datetime(2025,1,2,5,55)).accepted
    early_out=scan(engine,eid,datetime(2025,1,2,13,49))
    assert not early_out.accepted and early_out.reason.startswith('CLOCK_OUT_NOT_ALLOWED')
    assert scan(engine,eid,datetime(2025,1,2,13,50)).accepted
    with db.session() as s:
        session=s.scalar(__import__('sqlalchemy').select(AttendanceSession))
        assert session.clock_in==datetime(2025,1,2,5,55) and session.clock_out==datetime(2025,1,2,13,50)
        assert s.scalar(__import__('sqlalchemy').select(__import__('sqlalchemy').func.count()).select_from(ScanEvent))==3
        assert s.scalar(__import__('sqlalchemy').select(__import__('sqlalchemy').func.count()).select_from(__import__('app.database.models',fromlist=['ExceptionRecord']).ExceptionRecord))==1
    pytest.importorskip('openpyxl')
    from app.export import ExcelExporter
    from openpyxl import load_workbook
    workbook=load_workbook(ExcelExporter(db).export(tmp_path,date_from=datetime(2025,1,2).date(),date_to=datetime(2025,1,2).date()))
    assert workbook['Attendance Summary'].max_row==2 and workbook['Raw Scans'].max_row==4 and workbook['Exceptions'].max_row==2
def test_early_clock_in_rules(db):
    eid=employee(db); en=AttendanceEngine(db)
    assert scan(en,eid,datetime(2025,1,2,5,59)).accepted
    eid2=employee(db); assert not scan(en,eid2,datetime(2025,1,2,4,59)).accepted
def test_too_early_out_and_valid_out(db):
    eid=employee(db); en=AttendanceEngine(db); scan(en,eid,datetime(2025,1,2,6,0))
    assert not scan(en,eid,datetime(2025,1,2,13,49)).accepted
    assert scan(en,eid,datetime(2025,1,2,13,50)).accepted
def test_duplicate_scan_is_rejected(db):
    eid=employee(db); en=AttendanceEngine(db); assert scan(en,eid,datetime(2025,1,2,6,0)).accepted
    r=scan(en,eid,datetime(2025,1,2,6,0,3)); assert not r.accepted and r.reason=='DUPLICATE_SCAN'
def test_unknown_and_inactive_are_logged(db):
    en=AttendanceEngine(db); assert scan(en,'EMP999999',datetime(2025,1,2,6)).reason=='UNKNOWN_BARCODE'
    eid=employee(db); EmployeeService(db).update(eid,active=False); assert scan(en,eid,datetime(2025,1,2,6,1)).reason=='INACTIVE_EMPLOYEE'
    with db.session() as s: assert s.scalar(__import__('sqlalchemy').select(__import__('sqlalchemy').func.count()).select_from(ScanEvent))==2
def test_overnight_shift(db):
    eid=employee(db,2); en=AttendanceEngine(db); assert scan(en,eid,datetime(2025,1,2,21,30)).accepted
    assert scan(en,eid,datetime(2025,1,3,6,3)).accepted
def test_multiple_terminal_scanner_identity(db):
    eid=employee(db); en=AttendanceEngine(db); scan(en,eid,datetime(2025,1,2,6), 'Scanner 3')
    with db.session() as s: event=s.scalar(__import__('sqlalchemy').select(ScanEvent)); assert event.terminal_id==1 and event.scanner_id=='Scanner 3'
def test_auth_and_backup(db,tmp_path):
    auth=AuthService(db); auth.create_admin('admin','safe-password'); assert auth.authenticate('admin','safe-password')
    destination=backup_database(db.path,tmp_path/'backup'); assert destination.exists()
def test_restore_creates_safety_backup(db,tmp_path):
    from app.maintenance import restore_database
    baseline=backup_database(db.path,tmp_path/'source_backup')
    EmployeeService(db).create('Extra','Person')
    safety=restore_database(baseline,db.path)
    assert safety.exists()
    reloaded=Database(db.path)
    assert len(EmployeeService(reloaded).list())==0

def test_excel_export(db,tmp_path):
    pytest.importorskip('openpyxl')
    from app.export import ExcelExporter
    output=ExcelExporter(db).export(tmp_path)
    assert output.exists() and output.suffix=='.xlsx'

def test_excel_export_filters(db,tmp_path):
    pytest.importorskip('openpyxl')
    from app.export import ExcelExporter
    eid=employee(db); engine=AttendanceEngine(db); scan(engine,eid,datetime(2025,1,2,6)); scan(engine,eid,datetime(2025,1,2,14))
    output=ExcelExporter(db).export(tmp_path,date_from=datetime(2025,1,2).date(),date_to=datetime(2025,1,2).date(),employee_id=eid)
    from openpyxl import load_workbook
    assert load_workbook(output)['Attendance Summary'].max_row==2

def test_barcode_generation(tmp_path):
    pytest.importorskip('barcode')
    from app.barcode import generate_barcode
    assert generate_barcode('EMP000421',tmp_path).exists()
